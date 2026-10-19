"""cpi.json → 분석용 엑셀 (월이 행, 항목이 열)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HDR = PatternFill("solid", fgColor="1F3A4D")
HFONT = Font(color="FFFFFF", bold=True)


def _sheet(wb, title, months, cols, getter, fmt="0.00"):
    ws = wb.create_sheet(title)
    ws.append(["월"] + [c[1] for c in cols])
    ws.append(["코드"] + [c[0] for c in cols])
    for cell in ws[1] + ws[2]:
        cell.fill, cell.font = HDR, HFONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for k, ym in enumerate(months):
        ws.append([ym] + [getter(c[0], k) for c in cols])
    for row in ws.iter_rows(min_row=3, min_col=2):
        for cell in row:
            cell.number_format = fmt
    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 9
    for i in range(2, len(cols) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.row_dimensions[1].height = 42


def write(d, path):
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    lines = [
        ["미국 CPI 세부 데이터 (자동 생성)"],
        ["생성 시각(UTC)", d["generated_at"]], ["최신 지표월", d["latest"]],
        ["BLS 파일 갱신", d["bls_last_modified"]], [],
        ["전월비", "계절조정(SA). 계절조정 시리즈가 없는 항목(세입자보험·가사서비스·통신·기타개인서비스 등)은 원계열"],
        ["전년비", "원계열(NSA)"],
        ["3·6개월 연율", "SA 지수 기준 연율 환산"],
        ["가중치", "BLS 상대중요도(전체=100). 매년 1월 값을 기준으로 월별 가격 변화로 갱신"],
        ["기여도(전월비)", "전월 가중치 × 전월비 / 100 → 헤드라인 %p. 근원 기준은 근원 가중치로 나눔"],
        ["기여도(전년비)", "12개월 전 가중치 × 전년비 / 100 (근사)"],
        ["기타(잔여)", "발표치 − 항목 기여도 합"],
        ["일회성 z", "이번 달 전월비가 직전 12개월 평균에서 몇 표준편차 벗어났는지"],
        ["2025-10", "셧다운으로 CPI 미발표 → 빈칸"],
        ["출처", "BLS cu.data.0.Current, BLS 보도자료 표1·2(가중치), 클리블랜드·애틀랜타 연은, 트레이딩뷰 경제 캘린더"],
    ]
    for l in lines:
        readme.append(l)
    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 110
    readme["A1"].font = Font(bold=True, size=14)

    months = d["months"]
    meta = [m for m in d["meta"]]
    cols_all = [(m["code"], m["ko"]) for m in meta]
    items = [(m["code"], m["ko"]) for m in meta if m["type"] == "item"]
    S, C = d["series"], d["contrib"]

    info = wb.create_sheet("항목정보")
    info.append(["코드", "이름", "유형", "버킷", "상품/서비스", "계절조정", "BLS 영문명", "설명", "BLS 링크"])
    for m in meta:
        url = "" if m["type"] == "derived" or m.get("derived") else \
            f"https://data.bls.gov/timeseries/{'CUSR0000' if m.get('sa', True) else 'CUUR0000'}{m['code']}"
        info.append([m["code"], m["ko"], m["type"], m.get("bucket", ""), m.get("kind", ""),
                     "SA" if m.get("sa", True) else "NSA", m.get("en", ""), m.get("desc", ""), url])
    for cell in info[1]:
        cell.fill, cell.font = HDR, HFONT
    for col, w in zip("ABCDEFGHI", (12, 26, 9, 12, 10, 8, 40, 90, 50)):
        info.column_dimensions[col].width = w

    for title, key in (("전월비_SA", "mom"), ("전년비_NSA", "yoy"), ("3개월연율", "ann3"), ("6개월연율", "ann6")):
        _sheet(wb, title, months, cols_all, lambda c, k, key=key: S[c][key][k])
    _sheet(wb, "가중치", months, cols_all, lambda c, k: S[c]["w"][k], fmt="0.000")
    res = d["resid"]
    for title, key in (("기여도_헤드라인_MoM", "mom_h"), ("기여도_근원_MoM", "mom_c"),
                       ("기여도_헤드라인_YoY", "yoy_h"), ("기여도_근원_YoY", "yoy_c")):
        cols = items + [("RESID", "기타(잔여)")]
        _sheet(wb, title, months, cols,
               lambda c, k, key=key: res[key][k] if c == "RESID" else C[c][key][k], fmt="0.000")
    _sheet(wb, "일회성_z", months, cols_all, lambda c, k: d["flags"][c]["z"][k])
    ext = [(k, e["ko"]) for k, e in d["ext"].items()]
    for title, key in (("외부지표_MoM", "mom"), ("외부지표_YoY", "yoy")):
        _sheet(wb, title, months, ext, lambda c, k, key=key: d["ext"][c][key][k])
    _sheet(wb, "확산도", months, [("up", "상승 항목 비중(%)"), ("hot", "연율 3% 초과 비중(%)")],
           lambda c, k: d["diffusion"][c][k], fmt="0.0")

    cs = wb.create_sheet("컨센서스")
    cs.append(["지표월", "헤드라인 MoM 예상", "실제", "헤드라인 YoY 예상", "실제", "근원 MoM 예상", "실제", "근원 YoY 예상", "실제"])
    for ym in sorted(d["consensus"]):
        v = d["consensus"][ym]
        row = [ym]
        for k in ("h_mom", "h_yoy", "c_mom", "c_yoy"):
            row += [v.get(k, {}).get("f"), v.get(k, {}).get("a")]
        cs.append(row)
    fo = wb.create_sheet("FOMC")
    fo.append(["날짜", "결정 금리(상단, %)", "직전"])
    for f in d["fomc"]:
        fo.append([f["date"], f["rate"], f["prev"]])
    wb.save(path)
